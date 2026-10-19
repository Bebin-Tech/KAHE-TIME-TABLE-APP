from io import BytesIO
from django.test import TestCase
from accounts.models import User
from timetables.tests import SchedulingTests
from timetables.services import transition


class ReportTests(TestCase):
    def setUp(self):
        SchedulingTests.setUp(self)
        table, _ = SchedulingTests.save(self)
        for action in ["submit", "approve", "publish"]:
            table = transition(self.admin, table.pk, table.revision, action)

    def entry(self, **overrides):
        return SchedulingTests.entry(self, **overrides)

    def test_authorized_exports_are_real_files(self):
        from openpyxl import load_workbook

        self.client.force_login(self.student)
        response = self.client.get("/api/exports/pdf/")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(b"".join(response.streaming_content).startswith(b"%PDF"))
        response = self.client.get("/api/exports/xlsx/")
        book = load_workbook(BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(book.active.max_row, 2)
        self.assertEqual(book.active["F2"].value, "CS1 Programming")

    def test_foreign_section_and_anonymous_export(self):
        stranger = User.objects.create_user("stranger", role="STUDENT")
        self.client.force_login(stranger)
        self.assertEqual(
            self.client.get(f"/api/exports/pdf/?timetable={self.table.pk}").status_code,
            404,
        )
        self.assertEqual(self.client.get("/api/reports/").json()["rows"], [])
        self.client.logout()
        self.assertEqual(self.client.get("/api/exports/pdf/").status_code, 302)

    def test_published_master_changes_blocked(self):
        import json

        self.client.force_login(self.admin)
        data = {
            "code": "R1",
            "name": "Room 1",
            "room_type": "CLASSROOM",
            "capacity": 1,
            "equipment": [],
            "active": True,
        }
        result = self.client.put(
            f"/api/academic/rooms/{self.room.pk}/",
            json.dumps(data),
            content_type="application/json",
        )
        self.assertEqual(result.status_code, 400)
        self.room.refresh_from_db()
        self.assertEqual(self.room.capacity, 40)
