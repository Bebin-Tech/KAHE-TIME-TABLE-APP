from collections import Counter
from io import BytesIO
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import FileResponse, JsonResponse, Http404
from timetables.selectors import visible_timetables, visible_entries
from timetables.validators import interval
from academics.models import Period


def report_entries(request):
    tables = visible_timetables(request.user).filter(status="PUBLISHED")
    table_id = request.GET.get("timetable")
    if table_id:
        try:
            tables = tables.filter(pk=int(table_id))
        except ValueError:
            raise Http404
        if not tables.exists():
            raise Http404
    department = request.GET.get("department")
    if department:
        try:
            tables = tables.filter(
                section__semester__program__department_id=int(department)
            )
        except ValueError:
            raise Http404
    rows = []
    periods = list(Period.objects.all())
    days = [
        "Monday",
        "Tuesday",
        "Wednesday",
        "Thursday",
        "Friday",
        "Saturday",
        "Sunday",
    ]
    faculty_filter = request.GET.get("faculty")
    for table in tables:
        for e in visible_entries(table, request.user):
            if faculty_filter and str(e.faculty_id) != faculty_filter:
                continue
            times = interval(e, periods)
            rows.append(
                {
                    "section": str(table.section),
                    "version": table.version,
                    "day": days[e.day],
                    "starts": str(times[0]) if times else "",
                    "ends": str(times[1]) if times else "",
                    "subject": e.subject.code + " " + e.subject.name,
                    "faculty": str(e.faculty),
                    "room": e.room.code,
                    "periods": e.duration,
                    "required_subject_periods": e.subject.weekly_periods,
                }
            )
    return rows


@login_required
def reports(request):
    rows = report_entries(request)
    faculty = Counter()
    rooms = Counter()
    subjects = Counter()
    required = {}
    for r in rows:
        faculty[r["faculty"]] += r["periods"]
        rooms[r["room"]] += r["periods"]
        key = (r["section"], r["subject"])
        subjects[key] += r["periods"]
        required[key] = r["required_subject_periods"]
    result = [
        {
            "report": "Faculty workload",
            "resource": key,
            "scheduled_periods": count,
            "required_periods": "",
        }
        for key, count in faculty.items()
    ]
    if (
        request.user.role not in ["STUDENT", "FACULTY"]
        or request.user.is_institution_admin
    ):
        result += [
            {
                "report": "Room utilization",
                "resource": key,
                "scheduled_periods": count,
                "required_periods": "",
            }
            for key, count in rooms.items()
        ]
    if request.user.role not in ["FACULTY"] or request.user.is_institution_admin:
        from academics.models import Subject

        for table in visible_timetables(request.user).filter(status="PUBLISHED"):
            for subject in Subject.objects.filter(
                semester=table.section.semester, active=True
            ):
                key = (str(table.section), subject.code + " " + subject.name)
                required[key] = subject.weekly_periods
                subjects.setdefault(key, 0)
    result += [
        {
            "report": "Subject completion",
            "resource": key[0] + " · " + key[1],
            "scheduled_periods": count,
            "required_periods": required[key],
        }
        for key, count in subjects.items()
    ]
    return JsonResponse({"rows": result})


@login_required
def export(request, kind):
    rows = report_entries(request)
    buffer = BytesIO()
    if kind == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        book = Workbook()
        sheet = book.active
        sheet.title = "Published schedules"
        headers = [
            "section",
            "version",
            "day",
            "starts",
            "ends",
            "subject",
            "faculty",
            "room",
            "periods",
        ]
        sheet.append(headers)
        for row in rows:
            # Prefix formula-looking text so exports cannot execute spreadsheet formulas.
            sheet.append(
                [
                    (
                        "'" + row[k]
                        if isinstance(row[k], str)
                        and row[k].startswith(("=", "+", "-", "@"))
                        else row[k]
                    )
                    for k in headers
                ]
            )
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334EC9")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(
                55, max(len(str(c.value or "")) for c in col) + 3
            )
        summary = book.create_sheet("Workload and utilization")
        summary.append(["Metric", "Resource", "Weekly periods"])
        for field, label in [
            ("faculty", "Faculty workload"),
            ("room", "Room utilization"),
            ("subject", "Subject hours"),
        ]:
            counts = Counter()
            for row in rows:
                counts[row[field]] += row["periods"]
            for key, count in counts.items():
                summary.append(
                    [
                        label,
                        "'" + key if key.startswith(("=", "+", "-", "@")) else key,
                        count,
                    ]
                )
        book.save(buffer)
        content_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    elif kind == "pdf":
        from html import escape
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
        )

        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=24,
            leftMargin=24,
            topMargin=24,
            bottomMargin=24,
        )
        story = [
            Paragraph("College Timetable · Published schedules", styles["Title"]),
            Spacer(1, 14),
        ]
        headers = ["Section / version", "Day / time", "Subject", "Faculty", "Room"]
        data = [[Paragraph(h, styles["Heading5"]) for h in headers]]
        for row in rows:
            data.append(
                [
                    Paragraph(escape(str(v)), styles["BodyText"])
                    for v in [
                        row["section"] + f" / v{row['version']}",
                        row["day"] + " " + row["starts"][:5] + "–" + row["ends"][:5],
                        row["subject"],
                        row["faculty"],
                        row["room"],
                    ]
                ]
            )
        if not rows:
            story.append(
                Paragraph(
                    "No published schedules within your access.", styles["BodyText"]
                )
            )
        else:
            table = Table(data, colWidths=[200, 135, 200, 150, 100], repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9edfb")),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d8deeb")),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            story.append(table)
        doc.build(story)
        content_type = "application/pdf"
    else:
        raise Http404
    buffer.seek(0)
    return FileResponse(
        buffer,
        as_attachment=True,
        filename=f"college-timetable.{kind}",
        content_type=content_type,
    )
